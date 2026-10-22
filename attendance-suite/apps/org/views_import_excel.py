"""ورود مرحله‌ای پروندهٔ پرسنلی از Excel (ImportFromExcel)."""
from __future__ import annotations

import datetime as dt
import math
import os
from io import BytesIO
from zipfile import BadZipFile

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.core.services import audit, require_perm

from .models import Employee


IMPORT_FIELDS = {
    'badge_no': ('شماره کارگری', 'badgeNumberColumn'),
    'first_name': ('نام', 'firstNameColumn'),
    'last_name': ('نام خانوادگی', 'lastNameColumn'),
    'device_name': ('نام دستگاه', 'deviceNameColumn'),
    'email': ('ایمیل', 'emailColumn'),
    'tel': ('تلفن', 'telColumn'),
    'address': ('نشانی', 'addressColumn'),
    'personnel_no': ('شماره پرسنلی', 'numberColumn'),
    'gender': ('جنسیت', 'genderColumn'),
}
FIELD_LIMITS = {
    'badge_no': 24, 'first_name': 60, 'last_name': 80,
    'device_name': 120, 'email': 254, 'tel': 30,
    'address': 2000, 'personnel_no': 24, 'gender': 1,
}
SESSION_KEY = 't5_employee_import_preview'
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
MAX_COLUMNS = 200
PREVIEW_TTL_SECONDS = 30 * 60


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return '1' if value else '0'
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return ''
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def _gender(value):
    value = (value or '').strip().casefold()
    if not value:
        return ''
    if value in {'f', 'female', 'woman', 'زن', 'خانم', '1', 'true'}:
        return 'F'
    if value in {'m', 'male', 'man', 'مرد', 'آقا', '0', 'false'}:
        return 'M'
    return None


def _parse_column(value, max_column):
    try:
        column = int(value)
    except (TypeError, ValueError):
        raise ValueError('شمارهٔ ستون انتخاب‌شده معتبر نیست.')
    if not 1 <= column <= max_column:
        raise ValueError('شمارهٔ ستون باید در محدودهٔ فایل باشد.')
    return column - 1


def _build_rows(sheet, field_columns, from_row, to_row):
    raw_preview = []
    rows = []
    for row_number, row in enumerate(
            sheet.iter_rows(min_row=from_row, max_row=to_row,
                            max_col=min(sheet.max_column or 1, MAX_COLUMNS),
                            values_only=True), start=from_row):
        values = list(row)
        if len(raw_preview) < 30:
            raw_preview.append({
                'number': row_number,
                'cells': [_cell_text(value)[:300] for value in values[:20]],
            })
        payload = {}
        for field, column in field_columns.items():
            raw_value = _cell_text(values[column]) if column < len(values) else ''
            payload[field] = raw_value[:FIELD_LIMITS.get(field, 2000) + 1]
        if not any(payload.values()):
            continue
        rows.append({'row': row_number, 'payload': payload})
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValueError(f'حداکثر {MAX_IMPORT_ROWS} ردیف قابل بررسی است.')
    return rows, raw_preview


def _validate_candidates(rows):
    badges = [row['payload'].get('badge_no', '').strip()
              for row in rows if row['payload'].get('badge_no', '').strip()]
    existing = set(Employee.objects.filter(badge_no__in=badges)
                   .values_list('badge_no', flat=True))
    seen = set()
    candidates = []
    for row in rows:
        payload = {field: value.strip() for field, value in row['payload'].items()}
        payload.setdefault('device_name', '')
        payload.setdefault('email', '')
        payload.setdefault('tel', '')
        payload.setdefault('address', '')
        payload.setdefault('personnel_no', '')
        payload.setdefault('gender', '')
        badge = payload.get('badge_no', '')
        errors = []
        if not badge:
            errors.append('شماره کارگری وارد نشده است')
        if not payload.get('first_name'):
            errors.append('نام وارد نشده است')
        if not payload.get('last_name'):
            errors.append('نام خانوادگی وارد نشده است')
        if badge in seen and badge:
            errors.append('شماره کارگری در فایل تکراری است')
        if badge in existing:
            errors.append('شماره کارگری از قبل ثبت شده است')
        if badge:
            seen.add(badge)
        for field, maximum in FIELD_LIMITS.items():
            value = payload.get(field, '')
            if field == 'gender':
                normalized = _gender(value)
                if normalized is None:
                    errors.append('جنسیت قابل تشخیص نیست')
                elif normalized:
                    payload[field] = normalized
            elif value and len(value) > maximum:
                errors.append(f'{IMPORT_FIELDS[field][0]} بیش از حد مجاز است')
        if payload.get('email'):
            try:
                validate_email(payload['email'])
            except ValidationError:
                errors.append('نشانی ایمیل معتبر نیست')
        candidates.append({
            'row': row['row'], 'payload': payload,
            'importable': not errors,
            'status': 'قابل ورود' if not errors else '؛ '.join(errors),
        })
    return candidates


def _field_choices():
    required = {'badge_no', 'first_name', 'last_name'}
    return [{'key': key, 'label': value[0], 'column_name': value[1],
             'required': key in required}
            for key, value in IMPORT_FIELDS.items()]


@require_perm('org', 'manage')
def import_excel(request):
    """Upload -> preview with field mapping -> commit valid, nonduplicate employees."""
    if request.method == 'POST' and request.POST.get('action') == 'commit':
        preview = request.session.get(SESSION_KEY)
        if (not preview or preview.get('user_id') != request.user.pk
                or timezone.now().timestamp() - preview.get('created_at', 0)
                > PREVIEW_TTL_SECONDS):
            request.session.pop(SESSION_KEY, None)
            messages.error(request, 'پیش‌نمایش منقضی شده است؛ فایل را دوباره انتخاب کنید.')
            return redirect('org:import_excel')
        imported = skipped = 0
        with transaction.atomic():
            for item in preview['rows']:
                if not item['importable']:
                    continue
                try:
                    with transaction.atomic():
                        Employee.objects.create(**item['payload'])
                    imported += 1
                except IntegrityError:
                    skipped += 1
            audit(request.user, 'import_excel', 'Employee',
                  detail=f"file={preview['file_name']}; imported={imported}; skipped={skipped}")
        request.session.pop(SESSION_KEY, None)
        request.session.modified = True
        return render(request, 'org/import_excel.html', {
            'fields': _field_choices(),
            'result': {'imported': imported, 'skipped': skipped,
                       'file_name': preview['file_name']},
        })

    if request.method == 'POST' and request.POST.get('action') == 'preview':
        request.session.pop(SESSION_KEY, None)
        upload = request.FILES.get('file')
        if not upload:
            messages.error(request, 'فایل Excel را انتخاب کنید.')
            return redirect('org:import_excel')
        if upload.size > MAX_FILE_SIZE:
            messages.error(request, 'حداکثر اندازهٔ فایل ۱۰ مگابایت است.')
            return redirect('org:import_excel')
        if os.path.splitext(upload.name)[1].lower() not in {'.xlsx', '.xlsm'}:
            messages.error(request, 'فقط فایل Excel با پسوند xlsx یا xlsm پشتیبانی می‌شود.')
            return redirect('org:import_excel')

        try:
            workbook = load_workbook(BytesIO(upload.read()), read_only=True,
                                     data_only=True, keep_links=False)
            sheet = workbook.active
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError(f'فایل بیش از {MAX_COLUMNS} ستون دارد.')
            selected = request.POST.getlist('selected_fields')
            if any(field not in IMPORT_FIELDS for field in selected):
                raise ValueError('یکی از فیلدهای انتخاب‌شده معتبر نیست.')
            selected = list(dict.fromkeys(selected))
            required = {'badge_no', 'first_name', 'last_name'}
            if not required.issubset(selected):
                raise ValueError('شماره کارگری، نام و نام خانوادگی باید انتخاب شوند.')
            field_columns = {}
            for field in selected:
                field_columns[field] = _parse_column(
                    request.POST.get(f'column_{field}'), sheet.max_column or 0)
            if len(set(field_columns.values())) != len(field_columns):
                raise ValueError('هر فیلد باید به ستون جداگانه متصل شود.')
            try:
                from_row = int(request.POST.get('from_row') or 2)
                to_row = int(request.POST.get('to_row') or (sheet.max_row or 0))
            except ValueError:
                raise ValueError('شمارهٔ ردیف شروع/پایان معتبر نیست.')
            if (from_row < 1 or to_row < from_row or to_row > (sheet.max_row or 0)
                    or to_row - from_row + 1 > MAX_IMPORT_ROWS):
                raise ValueError(f'بازهٔ ردیف باید معتبر و حداکثر {MAX_IMPORT_ROWS} ردیف باشد.')
            source_rows, raw_preview = _build_rows(
                sheet, field_columns, from_row, to_row)
            candidates = _validate_candidates(source_rows)
            workbook.close()
        except (ValueError, TypeError, KeyError, OSError, InvalidFileException,
                BadZipFile) as exc:
            try:
                workbook.close()
            except (UnboundLocalError, AttributeError):
                pass
            messages.error(request, str(exc) or 'خواندن فایل Excel ممکن نشد.')
            return redirect('org:import_excel')

        request.session[SESSION_KEY] = {
            'user_id': request.user.pk,
            'created_at': timezone.now().timestamp(),
            'file_name': os.path.basename(upload.name)[:255],
            'rows': candidates,
            'raw_preview': raw_preview,
            'selected_fields': selected,
            'column_values': {key: value + 1 for key, value in field_columns.items()},
            'from_row': from_row,
            'to_row': to_row,
        }
        request.session.modified = True
        return render(request, 'org/import_excel.html', {
            'fields': _field_choices(),
            'preview': candidates,
            'raw_preview': raw_preview,
            'file_name': os.path.basename(upload.name),
            'valid_count': sum(row['importable'] for row in candidates),
            'invalid_count': sum(not row['importable'] for row in candidates),
            'selected_fields': selected,
            'column_values': {key: value + 1 for key, value in field_columns.items()},
            'from_row': from_row,
            'to_row': to_row,
        })

    preview = request.session.get(SESSION_KEY)
    if preview and (preview.get('user_id') != request.user.pk
                    or timezone.now().timestamp() - preview.get('created_at', 0)
                    > PREVIEW_TTL_SECONDS):
        request.session.pop(SESSION_KEY, None)
        request.session.modified = True
        preview = None
    if preview:
        return render(request, 'org/import_excel.html', {
            'fields': _field_choices(), 'preview': preview['rows'],
            'raw_preview': preview.get('raw_preview', []),
            'file_name': preview['file_name'],
            'valid_count': sum(row['importable'] for row in preview['rows']),
            'invalid_count': sum(not row['importable'] for row in preview['rows']),
            'selected_fields': preview.get('selected_fields', []),
            'column_values': preview.get('column_values', {}),
            'from_row': preview.get('from_row', 2),
            'to_row': preview.get('to_row', ''),
        })
    return render(request, 'org/import_excel.html', {'fields': _field_choices()})
