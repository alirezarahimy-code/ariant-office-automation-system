"""نماهای انبار حرفه‌ای: اسناد چندردیفی، کارت کالا، قیمت، انبارگردانی و گزارش‌ها.

الگوی نمایش مانند دبیرخانه: page-head/card/table/chip با تقویم شمسی و RTL.
"""
from __future__ import annotations

from decimal import Decimal

import jdatetime
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from apps.core.services import feature_enabled

from .models import (CountTag, InventoryPeriod, Item, ItemGroup, ItemSpec,
                     ItemUnit, PriceList, PriceListItem, SerialNumber,
                     StockCount, StockCountLine, StockDocument,
                     StockDocumentLine, StockMove, Warehouse,
                     WarehouseLocation)
from .services import (apply_batch_pricing, apply_count, cancel_document,
                       confirm_document, create_document,
                       create_return_document, current_stock,
                       expiry_report, generate_count_tags, import_count_rows,
                       inventory_value, issue, item_card, low_stock_items,
                       movement_report, no_movement_items, receive,
                       serial_report, set_item_price, transfer,
                       visible_warehouses)


def _can_manage(user) -> bool:
    return user.is_superuser or user.has_perm('inventory.manage_inventory')


def _can_view_value(user) -> bool:
    """تفکیک دسترسی تعدادی/ریالی — ارزش موجودی فقط با مجوز خاص."""
    return (user.is_superuser or _can_manage(user)
            or user.has_perm('inventory.view_inventory_value'))


def _parse_jalali(raw: str):
    from apps.secretariat.views import _parse_jalali_input
    return _parse_jalali_input(raw)


def _warehouses(request):
    return visible_warehouses(request.user)


# ─────────────────────────────────────────────────────────────
# کالا و کارت کالا
# ─────────────────────────────────────────────────────────────
@login_required
def item_list(request):
    q = request.GET.get('q', '').strip()
    group_id = request.GET.get('group', '')
    show_inactive = request.GET.get('inactive') == '1'
    items = Item.objects.all()
    if not show_inactive:
        items = items.filter(is_active=True)
    if q:
        items = items.filter(
            Q(name__icontains=q) | Q(code__icontains=q) | Q(category__icontains=q)
            | Q(barcode__icontains=q))
    groups = ItemGroup.objects.all()
    if group_id:
        ids = [int(group_id)]
        group = ItemGroup.objects.filter(pk=group_id).first()
        if group is not None:
            ids = group.get_descendant_ids()
        items = items.filter(group_id__in=ids)
    rows = [{'item': item, 'stock': current_stock(item)} for item in items]
    return render(request, 'inventory/list.html', {
        'rows': rows, 'q': q, 'groups': groups, 'group_id': group_id,
        'show_inactive': show_inactive,
        'low': low_stock_items(), 'can_manage': _can_manage(request.user),
    })


@login_required
def item_detail(request, pk):
    item = get_object_or_404(Item.objects.select_related('group', 'default_location'),
                             pk=pk)
    rows = item_card(item)
    return render(request, 'inventory/detail.html', {
        'item': item, 'rows': rows, 'stock': current_stock(item),
        'warehouses': _warehouses(request),
        'show_value': _can_view_value(request.user),
        'specs': item.specs.all(), 'units': item.units.all(),
        'serials': item.serials.all() if feature_enabled('inventory.serial_expiry') else [],
        'prices': item.price_entries.select_related('price_list')
        if feature_enabled('inventory.price_lists') else [],
    })


def _save_item_relations(item, post):
    """ذخیره واحدها و مشخصات از ردیف‌های فرم (نام‌گذاری ایندکس‌دار)."""
    for i in range(25):
        name = post.get(f'unit_name_{i}', '').strip()
        if not name:
            continue
        try:
            ratio = Decimal(post.get(f'unit_ratio_{i}', '1') or '1')
        except ArithmeticError:
            ratio = Decimal('1')
        is_default = post.get(f'unit_default_{i}') == 'on'
        ItemUnit.objects.update_or_create(
            item=item, name=name,
            defaults={'ratio': ratio, 'is_default': is_default})
    for i in range(25):
        title = post.get(f'spec_title_{i}', '').strip()
        value = post.get(f'spec_value_{i}', '').strip()
        if not title:
            continue
        ItemSpec.objects.create(item=item, title=title, value=value)


def _item_form_context(item=None):
    return {
        'item': item,
        'groups': ItemGroup.objects.all(),
        'locations': WarehouseLocation.objects.all(),
        'costing_methods': Item.CostingMethod.choices,
    }


@login_required
def item_new(request):
    if not _can_manage(request.user):
        messages.error(request, 'امکان ثبت کالا ندارید.')
        return redirect('inventory:list')
    if request.method != 'POST':
        return render(request, 'inventory/item_form.html', _item_form_context())
    code = request.POST.get('code', '').strip() or Item.next_code()
    item = Item.objects.create(
        code=code, name=request.POST.get('name', '').strip(),
        unit=request.POST.get('unit', '').strip() or 'عدد',
        category=request.POST.get('category', '').strip(),
        group_id=request.POST.get('group') or None,
        description=request.POST.get('description', '').strip(),
        barcode=request.POST.get('barcode', '').strip(),
        is_service=request.POST.get('is_service') == 'on',
        is_active=request.POST.get('is_active', 'on') == 'on',
        min_stock=Decimal(request.POST.get('min_stock', '0') or '0'),
        max_stock=Decimal(request.POST.get('max_stock', '0') or '0'),
        reorder_point=Decimal(request.POST.get('reorder_point', '0') or '0'),
        track_serial=request.POST.get('track_serial') == 'on',
        has_expiry=request.POST.get('has_expiry') == 'on',
        costing_method=request.POST.get('costing_method', Item.CostingMethod.AVERAGE),
        default_location_id=request.POST.get('default_location') or None)
    if request.FILES.get('image'):
        item.image = request.FILES['image']
        item.save(update_fields=['image', 'updated_at'])
    _save_item_relations(item, request.POST)
    messages.success(request, f'کالای {item.code} ثبت شد.')
    return redirect('inventory:detail', pk=item.pk)


@login_required
def item_edit(request, pk):
    if not _can_manage(request.user):
        messages.error(request, 'امکان ویرایش کالا ندارید.')
        return redirect('inventory:detail', pk=pk)
    item = get_object_or_404(Item, pk=pk)
    if request.method != 'POST':
        return render(request, 'inventory/item_form.html', _item_form_context(item))
    item.name = request.POST.get('name', '').strip()
    item.unit = request.POST.get('unit', '').strip() or 'عدد'
    item.category = request.POST.get('category', '').strip()
    item.group_id = request.POST.get('group') or None
    item.description = request.POST.get('description', '').strip()
    item.barcode = request.POST.get('barcode', '').strip()
    item.is_service = request.POST.get('is_service') == 'on'
    item.is_active = request.POST.get('is_active') == 'on'
    item.min_stock = Decimal(request.POST.get('min_stock', '0') or '0')
    item.max_stock = Decimal(request.POST.get('max_stock', '0') or '0')
    item.reorder_point = Decimal(request.POST.get('reorder_point', '0') or '0')
    item.track_serial = request.POST.get('track_serial') == 'on'
    item.has_expiry = request.POST.get('has_expiry') == 'on'
    item.costing_method = request.POST.get('costing_method', Item.CostingMethod.AVERAGE)
    item.default_location_id = request.POST.get('default_location') or None
    if request.FILES.get('image'):
        item.image = request.FILES['image']
    item.save()
    item.units.all().delete()
    item.specs.all().delete()
    _save_item_relations(item, request.POST)
    messages.success(request, f'کالای {item.code} به‌روزرسانی شد.')
    return redirect('inventory:detail', pk=item.pk)


@login_required
def group_list(request):
    if not feature_enabled('inventory.item_groups'):
        return render(request, 'core/feature_disabled.html', {'feature_code': 'inventory.item_groups'})
    if request.method == 'POST' and _can_manage(request.user):
        parent_id = request.POST.get('parent') or None
        group = ItemGroup.objects.create(
            name=request.POST.get('name', '').strip(),
            parent_id=parent_id)
        messages.success(request, f'گروه «{group.name}» ثبت شد.')
        return redirect('inventory:groups')
    groups = ItemGroup.objects.all()
    return render(request, 'inventory/group_list.html', {
        'groups': groups, 'can_manage': _can_manage(request.user),
    })


# ─────────────────────────────────────────────────────────────
# اسناد انبار (چندردیفی)
# ─────────────────────────────────────────────────────────────
def _parse_doc_lines(post, doc_type):
    """استخراج ردیف‌های سند از فرم (فیلدهای ایندکس‌دار)."""
    lines = []
    i = 0
    while i < 30:
        item_id = post.get(f'line_item_{i}', '').strip()
        if not item_id:
            i += 1
            if i > 3 and not any(post.get(f'line_item_{j}', '').strip()
                                 for j in range(i, i + 3)):
                break
            continue
        item = Item.objects.filter(pk=item_id).first()
        if item is None:
            raise ValueError(f'کالای ردیف {i + 1} نامعتبر است.')
        quantity = Decimal(post.get(f'line_quantity_{i}', '0') or '0')
        unit_price_raw = post.get(f'line_price_{i}', '').strip().replace(',', '')
        unit_price = Decimal(unit_price_raw) if unit_price_raw else None
        expiry = _parse_jalali(post.get(f'line_expiry_{i}', ''))
        lines.append({
            'item': item, 'quantity': quantity, 'unit_price': unit_price,
            'unit_name': post.get(f'line_unit_{i}', '').strip(),
            'batch_no': post.get(f'line_batch_{i}', '').strip(),
            'expiry_date': expiry,
            'serials_raw': post.get(f'line_serials_{i}', ''),
            'description': post.get(f'line_desc_{i}', '').strip(),
        })
        i += 1
    return lines


@login_required
def doc_list(request):
    if not feature_enabled('inventory.documents'):
        return render(request, 'core/feature_disabled.html', {'feature_code': 'inventory.documents'})
    docs = StockDocument.objects.select_related(
        'warehouse', 'to_warehouse', 'created_by')
    doc_type = request.GET.get('type', '')
    status = request.GET.get('status', '')
    warehouse_id = request.GET.get('warehouse', '')
    q = request.GET.get('q', '').strip()
    date_from = _parse_jalali(request.GET.get('date_from', ''))
    date_to = _parse_jalali(request.GET.get('date_to', ''))
    if doc_type in StockDocument.DocType.values:
        docs = docs.filter(doc_type=doc_type)
    if status in StockDocument.Status.values:
        docs = docs.filter(status=status)
    if warehouse_id:
        docs = docs.filter(Q(warehouse_id=warehouse_id) | Q(to_warehouse_id=warehouse_id))
    if q:
        docs = docs.filter(
            Q(number__icontains=q) | Q(reference__icontains=q)
            | Q(description__icontains=q))
    if date_from:
        docs = docs.filter(date__gte=date_from.togregorian())
    if date_to:
        docs = docs.filter(date__lte=date_to.togregorian())
    return render(request, 'inventory/doc_list.html', {
        'docs': docs[:200], 'warehouses': _warehouses(request),
        'types': StockDocument.DocType.choices, 'statuses': StockDocument.Status.choices,
        'cur': {'type': doc_type, 'status': status, 'warehouse': warehouse_id,
                'q': q, 'date_from': request.GET.get('date_from', ''),
                'date_to': request.GET.get('date_to', '')},
        'can_manage': _can_manage(request.user),
    })


@login_required
def doc_new(request):
    if not feature_enabled('inventory.documents'):
        return render(request, 'core/feature_disabled.html', {'feature_code': 'inventory.documents'})
    doc_type = request.GET.get('type', StockDocument.DocType.RECEIVE)
    if doc_type not in StockDocument.DocType.values:
        doc_type = StockDocument.DocType.RECEIVE
    if request.method != 'POST':
        return render(request, 'inventory/doc_form.html', {
            'items': Item.objects.filter(is_active=True),
            'warehouses': _warehouses(request),
            'locations': WarehouseLocation.objects.all(),
            'doc_type': doc_type,
            'types': StockDocument.DocType.choices,
            'categories': StockMove.Category.choices,
            'auto_finalize': not feature_enabled('inventory.doc_temp'),
        })
    post_type = request.POST.get('doc_type', doc_type)
    if post_type not in StockDocument.DocType.values:
        post_type = StockDocument.DocType.RECEIVE
    jdate = _parse_jalali(request.POST.get('date', '')) or jdatetime.date.today()
    category = request.POST.get('category', StockMove.Category.OTHER)
    if category not in StockMove.Category.values:
        category = StockMove.Category.OTHER
    try:
        lines = _parse_doc_lines(request.POST, post_type)
        if not lines:
            raise ValueError('سند حداقل یک ردیف لازم دارد.')
        to_warehouse = None
        if post_type == StockDocument.DocType.TRANSFER:
            to_warehouse = get_object_or_404(
                Warehouse, pk=request.POST.get('to_warehouse_id'))
        finalize = (request.POST.get('finalize') == '1'
                    or not feature_enabled('inventory.doc_temp'))
        doc = create_document(
            doc_type=post_type,
            warehouse=get_object_or_404(Warehouse, pk=request.POST.get('warehouse_id')),
            to_warehouse=to_warehouse, date=jdate, created_by=request.user,
            category=category, reference=request.POST.get('reference', '').strip(),
            description=request.POST.get('description', '').strip(),
            finalize=False, lines=[])
        for i, line_data in enumerate(lines):
            from .services import add_line
            add_line(doc, sequence=i, **line_data)
        if finalize:
            confirm_document(doc, confirmed_by=request.user)
    except (ValueError, ArithmeticError) as exc:
        messages.error(request, str(exc))
        return redirect('inventory:doc_new')
    if doc.status == StockDocument.Status.FINAL:
        messages.success(request, f'سند {doc.number} قطعی شد.')
    else:
        messages.success(request,
                         f'سند {doc.number} به‌صورت موقت ثبت شد؛ پس از بازبینی قطعی کنید.')
    return redirect('inventory:doc_detail', pk=doc.pk)


@login_required
def doc_detail(request, pk):
    if not feature_enabled('inventory.documents'):
        return render(request, 'core/feature_disabled.html', {'feature_code': 'inventory.documents'})
    doc = get_object_or_404(
        StockDocument.objects.select_related(
            'warehouse', 'to_warehouse', 'created_by', 'confirmed_by',
            'cancelled_by', 'reverses'),
        pk=pk)
    return render(request, 'inventory/doc_detail.html', {
        'doc': doc, 'lines': doc.lines.select_related('item', 'location'),
        'moves': doc.moves.select_related('item', 'warehouse'),
        'can_manage': _can_manage(request.user),
    })


@login_required
@require_POST
def doc_confirm(request, pk):
    doc = get_object_or_404(StockDocument, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان قطعی کردن سند ندارید.')
        return redirect('inventory:doc_detail', pk=pk)
    try:
        confirm_document(doc, confirmed_by=request.user)
        messages.success(request, f'سند {doc.number} قطعی شد.')
    except ValueError as exc:
        messages.error(request, str(exc))
    return redirect('inventory:doc_detail', pk=pk)


@login_required
@require_POST
def doc_cancel(request, pk):
    doc = get_object_or_404(StockDocument, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان ابطال سند ندارید.')
        return redirect('inventory:doc_detail', pk=pk)
    try:
        cancel_document(doc, reason=request.POST.get('reason', ''),
                        cancelled_by=request.user)
        messages.success(request, f'سند {doc.number} ابطال شد.')
    except ValueError as exc:
        messages.error(request, str(exc))
    return redirect('inventory:doc_detail', pk=pk)


@login_required
@require_POST
def doc_return(request, pk):
    doc = get_object_or_404(StockDocument, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان صدور سند برگشتی ندارید.')
        return redirect('inventory:doc_detail', pk=pk)
    try:
        ret = create_return_document(doc, created_by=request.user)
        messages.success(request, f'سند برگشتی {ret.number} صادر شد.')
        return redirect('inventory:doc_detail', pk=ret.pk)
    except ValueError as exc:
        messages.error(request, str(exc))
    return redirect('inventory:doc_detail', pk=pk)


@login_required
def doc_print(request, pk):
    """چاپ سند چندردیفی — سربرگ‌دار با شماره سند و امضای تحویل‌گیرنده."""
    doc = get_object_or_404(
        StockDocument.objects.select_related('warehouse', 'to_warehouse',
                                             'created_by', 'confirmed_by'),
        pk=pk)
    return render(request, 'inventory/doc_print.html', {
        'doc': doc, 'lines': doc.lines.select_related('item'),
    })


# ─────────────────────────────────────────────────────────────
# سازگاری: فرم تک‌قلمی قدیمی (رسید/حواله)
# ─────────────────────────────────────────────────────────────
@login_required
def move_new(request):
    if request.method != 'POST':
        return render(request, 'inventory/move_new.html', {
            'items': Item.objects.filter(is_active=True),
            'warehouses': _warehouses(request),
        })
    item = get_object_or_404(Item, pk=request.POST.get('item_id'))
    warehouse = get_object_or_404(Warehouse, pk=request.POST.get('warehouse_id'))
    kind = request.POST.get('kind', StockMove.Kind.RECEIVE)
    category = request.POST.get('category', StockMove.Category.OTHER)
    if category not in StockMove.Category.values:
        category = StockMove.Category.OTHER
    jdate = _parse_jalali(request.POST.get('date', '')) or jdatetime.date.today()
    try:
        quantity = Decimal(request.POST.get('quantity', '0').replace(',', ''))
        unit_price = request.POST.get('unit_price', '').strip().replace(',', '')
        unit_price = Decimal(unit_price) if unit_price else None
        if kind == StockMove.Kind.ISSUE:
            move = issue(item=item, warehouse=warehouse, quantity=quantity,
                         date=jdate, reference=request.POST.get('reference', ''),
                         note=request.POST.get('note', ''), created_by=request.user,
                         category=category, unit_price=unit_price)
        elif kind == StockMove.Kind.RECEIVE:
            move = receive(item=item, warehouse=warehouse, quantity=quantity,
                           date=jdate, reference=request.POST.get('reference', ''),
                           note=request.POST.get('note', ''), created_by=request.user,
                           category=category, unit_price=unit_price)
        else:
            raise ValueError('نوع سند نامعتبر است.')
    except (ValueError, ArithmeticError) as exc:
        messages.error(request, str(exc))
        return redirect('inventory:move_new')
    messages.success(request, f'{move.get_kind_display()} {item.code} ثبت شد.')
    return redirect('inventory:detail', pk=item.pk)


@login_required
@require_POST
def transfer_new(request):
    """انتقال کالا بین انبارها — سند انتقال با حواله مبدأ + رسید مقصد."""
    item = get_object_or_404(Item, pk=request.POST.get('item_id'))
    from_wh = get_object_or_404(Warehouse, pk=request.POST.get('from_warehouse_id'))
    to_wh = get_object_or_404(Warehouse, pk=request.POST.get('to_warehouse_id'))
    jdate = _parse_jalali(request.POST.get('date', '')) or jdatetime.date.today()
    try:
        quantity = Decimal(
            request.POST.get('quantity', '0').replace(',', ''))
        transfer(item=item, from_warehouse=from_wh, to_warehouse=to_wh,
                 quantity=quantity, date=jdate,
                 reference=request.POST.get('reference', ''),
                 note=request.POST.get('note', ''), created_by=request.user)
    except (ValueError, ArithmeticError) as exc:
        messages.error(request, str(exc))
        return redirect('inventory:move_new')
    messages.success(request,
                     f'انتقال {item.code} از {from_wh} به {to_wh} ثبت شد.')
    return redirect('inventory:detail', pk=item.pk)


@login_required
def move_print(request, pk):
    """چاپ رویداد تک‌قلمی (سازگاری قدیمی)."""
    move = get_object_or_404(
        StockMove.objects.select_related('item', 'warehouse', 'created_by'), pk=pk)
    return render(request, 'inventory/move_print.html', {'move': move})


# ─────────────────────────────────────────────────────────────
# موجودی، ارزش و ریزگردش
# ─────────────────────────────────────────────────────────────
@login_required
def balance(request):
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    rows = [{'item': item, 'stock': current_stock(item, warehouse)} for item in
            Item.objects.filter(is_active=True)]
    return render(request, 'inventory/balance.html', {
        'rows': rows, 'low': low_stock_items(warehouse),
        'warehouses': _warehouses(request), 'warehouse': warehouse,
    })


@login_required
def value_report(request):
    """ارزش موجودی انبار — میانگین موزون یا FIFO."""
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    data = inventory_value(warehouse)
    show_value = _can_view_value(request.user)
    return render(request, 'inventory/value.html', {
        'data': data, 'warehouses': _warehouses(request),
        'warehouse': warehouse, 'show_value': show_value,
    })


@login_required
def movement(request, pk):
    """ریزگردش کالا در بازه شمسی — با موجودی ابتدای بازه و پیوسته."""
    item = get_object_or_404(Item, pk=pk)
    date_from = _parse_jalali(request.GET.get('date_from', ''))
    date_to = _parse_jalali(request.GET.get('date_to', ''))
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    report = movement_report(item, date_from, date_to, warehouse)
    return render(request, 'inventory/movement.html', {
        'item': item, 'report': report,
        'warehouses': _warehouses(request),
        'date_from': request.GET.get('date_from', ''),
        'date_to': request.GET.get('date_to', ''),
        'warehouse': warehouse,
    })


# ─────────────────────────────────────────────────────────────
# انبارگردانی حرفه‌ای
# ─────────────────────────────────────────────────────────────
@login_required
def count_list(request):
    counts = StockCount.objects.select_related('warehouse', 'created_by')
    return render(request, 'inventory/count_list.html', {
        'counts': counts, 'can_manage': _can_manage(request.user),
        'warehouses': _warehouses(request),
    })


@login_required
@require_POST
def count_new(request):
    if not _can_manage(request.user):
        messages.error(request, 'امکان انبارگردانی ندارید.')
        return redirect('inventory:count_list')
    warehouse = get_object_or_404(Warehouse, pk=request.POST.get('warehouse_id'))
    jdate = _parse_jalali(request.POST.get('date', '')) or jdatetime.date.today()
    count = StockCount.objects.create(
        warehouse=warehouse, date=jdate, created_by=request.user,
        title=request.POST.get('title', '').strip())
    count.status = StockCount.Status.COUNTING
    count.save(update_fields=['status', 'updated_at'])
    if feature_enabled('inventory.count_tags'):
        generate_count_tags(count, quantity=int(
            request.POST.get('tag_count', '10') or '10'))
    messages.success(request, f'انبارگردانی {warehouse} آغاز شد.')
    return redirect('inventory:count_detail', pk=count.pk)


@login_required
def count_detail(request, pk):
    count = get_object_or_404(
        StockCount.objects.select_related('warehouse', 'created_by'), pk=pk)
    lines = []
    for line in count.lines.select_related('item'):
        expected = current_stock(line.item, count.warehouse)
        lines.append({
            'line': line, 'expected': expected,
            'diff': line.counted_quantity - expected
            if line.counted_quantity is not None else None})
    return render(request, 'inventory/count_detail.html', {
        'count': count, 'lines': lines, 'tags': count.tags.all()[:50],
        'items': Item.objects.filter(is_active=True),
        'can_manage': _can_manage(request.user),
    })


@login_required
@require_POST
def count_line_save(request, pk):
    """ثبت/به‌روزرسانی شمارش یک کالا (مرحله ۱ یا ۲)."""
    count = get_object_or_404(StockCount, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان ثبت شمارش ندارید.')
        return redirect('inventory:count_detail', pk=pk)
    if count.status == StockCount.Status.APPLIED:
        messages.error(request, 'این انبارگردانی اعمال شده است.')
        return redirect('inventory:count_detail', pk=pk)
    item = get_object_or_404(Item, pk=request.POST.get('item_id'))
    stage = int(request.POST.get('stage', '1') or '1')
    raw = request.POST.get('quantity', '').strip().replace(',', '')
    if not raw:
        messages.error(request, 'تعداد شمارش را وارد کنید.')
        return redirect('inventory:count_detail', pk=pk)
    qty = Decimal(raw)
    line, _ = StockCountLine.objects.get_or_create(
        count=count, item=item, defaults={'counted_quantity': qty})
    if stage == 1:
        line.stage1_quantity = qty
    else:
        line.stage2_quantity = qty
    line.counted_quantity = qty
    line.note = request.POST.get('note', line.note)
    line.save()
    messages.success(request, f'شمارش {item.code} ثبت شد.')
    return redirect('inventory:count_detail', pk=pk)


@login_required
@require_POST
def count_import(request, pk):
    """ورود شمارش از فایل اکسل (کد کالا / تعداد)."""
    count = get_object_or_404(StockCount, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان ورود شمارش ندارید.')
        return redirect('inventory:count_detail', pk=pk)
    upload = request.FILES.get('file')
    if upload is None:
        messages.error(request, 'فایل اکسل را انتخاب کنید.')
        return redirect('inventory:count_detail', pk=pk)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            code = str(row[0]).strip()
            item = Item.objects.filter(Q(code=code) | Q(barcode=code)).first()
            if item is None:
                continue
            rows.append((item, row[1]))
        stage = int(request.POST.get('stage', '1') or '1')
        updated = import_count_rows(count, rows, stage=stage)
        messages.success(request, f'{updated} ردیف شمارش وارد شد.')
    except Exception as exc:  # noqa: BLE001
        messages.error(request, f'خطای خواندن فایل: {exc}')
    return redirect('inventory:count_detail', pk=pk)


@login_required
@require_POST
def count_tags(request, pk):
    """صدور تگ‌های چاپی انبارگردانی."""
    count = get_object_or_404(StockCount, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان صدور تگ ندارید.')
        return redirect('inventory:count_detail', pk=pk)
    try:
        quantity = int(request.POST.get('quantity', '10') or '10')
        generate_count_tags(count, quantity=quantity)
        messages.success(request, f'{quantity} تگ صدور شد.')
    except ValueError as exc:
        messages.error(request, str(exc))
    return redirect('inventory:count_tags_print', pk=count.pk)


@login_required
def count_tags_print(request, pk):
    """چاپ تگ‌های انبارگردانی — برچسب بارکد‌دار."""
    count = get_object_or_404(
        StockCount.objects.select_related('warehouse'), pk=pk)
    tags = count.tags.select_related('item').all()
    from .barcode import code39_svg
    return render(request, 'inventory/count_tags.html', {
        'count': count,
        'tags': [{'tag': t, 'svg': code39_svg(t.code)} for t in tags],
    })


@login_required
@require_POST
def count_apply_view(request, pk):
    """اعمال انبارگردانی — سند اصلاحی خودکار (کسری/اضافی)."""
    count = get_object_or_404(StockCount, pk=pk)
    if not _can_manage(request.user):
        messages.error(request, 'امکان انبارگردانی ندارید.')
        return redirect('inventory:count_detail', pk=pk)
    try:
        apply_count(count, created_by=request.user)
        messages.success(request, f'انبارگردانی {count.warehouse} اعمال شد.')
        return redirect('inventory:count_detail', pk=count.pk)
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect('inventory:count_detail', pk=count.pk)


@login_required
@require_POST
def count_apply(request):
    """سازگاری: فرم قدیمی اعمال انبارگردانی از صفحه موجودی."""
    if not _can_manage(request.user):
        messages.error(request, 'امکان انبارگردانی ندارید.')
        return redirect('inventory:balance')
    warehouse = get_object_or_404(Warehouse, pk=request.POST.get('warehouse_id'))
    jdate = _parse_jalali(request.POST.get('date', '')) or jdatetime.date.today()
    count = StockCount.objects.create(
        warehouse=warehouse, date=jdate, created_by=request.user)
    line_count = 0
    for item in Item.objects.all():
        raw = request.POST.get(f'count_{item.pk}', '').strip()
        if not raw:
            continue
        try:
            qty = Decimal(raw.replace(',', ''))
        except ArithmeticError:
            continue
        StockCountLine.objects.create(
            count=count, item=item, counted_quantity=qty)
        line_count += 1
    if line_count == 0:
        count.delete()
        messages.error(request, 'هیچ شمارشی ثبت نشد.')
        return redirect('inventory:balance')
    apply_count(count, created_by=request.user)
    messages.success(request,
                     f'انبارگردانی {warehouse} با {line_count} ردیف اعمال شد.')
    return redirect('inventory:balance')


# ─────────────────────────────────────────────────────────────
# قیمت‌گذاری
# ─────────────────────────────────────────────────────────────
@login_required
def price_lists(request):
    if not feature_enabled('inventory.price_lists'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.price_lists'})
    if request.method == 'POST' and _can_manage(request.user):
        price_list = PriceList.objects.create(
            name=request.POST.get('name', '').strip(),
            is_default=request.POST.get('is_default') == 'on')
        messages.success(request, f'لیست قیمت «{price_list.name}» ثبت شد.')
        return redirect('inventory:price_list_detail', pk=price_list.pk)
    return render(request, 'inventory/price_lists.html', {
        'lists': PriceList.objects.all(),
        'can_manage': _can_manage(request.user),
    })


@login_required
def price_list_detail(request, pk):
    if not feature_enabled('inventory.price_lists'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.price_lists'})
    price_list = get_object_or_404(PriceList, pk=pk)
    if request.method == 'POST' and _can_manage(request.user):
        action = request.POST.get('action', 'set')
        try:
            if action == 'batch':
                percent = float(request.POST.get('percent', '0') or '0')
                count = apply_batch_pricing(price_list, percent)
                messages.success(request,
                                 f'قیمت {count} کالا به‌صورت دسته‌ای تغییر کرد.')
            else:
                item = get_object_or_404(Item, pk=request.POST.get('item_id'))
                set_item_price(price_list, item, request.POST.get('price', '0'))
                messages.success(request, f'قیمت {item.code} ثبت شد.')
        except (ValueError, ArithmeticError) as exc:
            messages.error(request, str(exc))
        return redirect('inventory:price_list_detail', pk=pk)
    entries = {e.item_id: e for e in price_list.entries.all()}
    rows = [{'item': item, 'entry': entries.get(item.pk)}
            for item in Item.objects.filter(is_active=True)]
    return render(request, 'inventory/price_list_detail.html', {
        'price_list': price_list, 'rows': rows,
        'can_manage': _can_manage(request.user),
    })


# ─────────────────────────────────────────────────────────────
# گزارش‌های تکمیلی
# ─────────────────────────────────────────────────────────────
@login_required
def report_no_movement(request):
    if not feature_enabled('inventory.no_movement'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.no_movement'})
    date_from = _parse_jalali(request.GET.get('date_from', ''))
    date_to = _parse_jalali(request.GET.get('date_to', ''))
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    rows = no_movement_items(date_from, date_to, warehouse)
    return render(request, 'inventory/report_no_movement.html', {
        'rows': rows, 'warehouses': _warehouses(request), 'warehouse': warehouse,
        'date_from': request.GET.get('date_from', ''),
        'date_to': request.GET.get('date_to', ''),
    })


@login_required
def report_expiry(request):
    if not feature_enabled('inventory.serial_expiry'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.serial_expiry'})
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    rows = expiry_report(warehouse)
    return render(request, 'inventory/report_expiry.html', {
        'rows': rows, 'warehouses': _warehouses(request), 'warehouse': warehouse,
    })


@login_required
def report_serials(request):
    if not feature_enabled('inventory.serial_expiry'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.serial_expiry'})
    item = None
    if request.GET.get('item'):
        item = Item.objects.filter(pk=request.GET.get('item')).first()
    return render(request, 'inventory/report_serials.html', {
        'rows': serial_report(item), 'item': item,
        'items': Item.objects.filter(track_serial=True),
    })


@login_required
def barcode_labels(request):
    """چاپ برچسب بارکد کالاها."""
    if not feature_enabled('inventory.barcode_labels'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.barcode_labels'})
    from .barcode import code39_svg
    ids = request.GET.getlist('item')
    items = Item.objects.filter(pk__in=ids) if ids else Item.objects.filter(
        is_active=True).exclude(barcode='')[:60]
    labels = []
    for item in items:
        value = item.barcode or item.code
        labels.append({'item': item, 'value': value,
                       'svg': code39_svg(value, bar_width=1.3, bar_height=32)})
    return render(request, 'inventory/barcode_labels.html', {
        'labels': labels, 'items': Item.objects.filter(is_active=True),
    })


# ─────────────────────────────────────────────────────────────
# انبارها، دوره مالی و خروجی‌ها
# ─────────────────────────────────────────────────────────────
@login_required
def warehouse_list(request):
    if request.method == 'POST' and _can_manage(request.user):
        action = request.POST.get('action', 'create')
        if action == 'users':
            warehouse = get_object_or_404(Warehouse, pk=request.POST.get('warehouse_id'))
            warehouse.users.set(request.POST.getlist('users'))
            messages.success(request, f'انبارداران {warehouse} به‌روزرسانی شد.')
        else:
            warehouse = Warehouse.objects.create(
                name=request.POST.get('name', '').strip(),
                description=request.POST.get('description', '').strip())
            messages.success(request, f'انبار «{warehouse.name}» ثبت شد.')
        return redirect('inventory:warehouses')
    warehouses = Warehouse.objects.prefetch_related('users', 'locations')
    users = User.objects.filter(is_active=True)
    return render(request, 'inventory/warehouse_list.html', {
        'warehouses': warehouses, 'users': users,
        'can_manage': _can_manage(request.user),
    })


@login_required
def period_list(request):
    if not feature_enabled('inventory.financial_period'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.financial_period'})
    if request.method == 'POST' and _can_manage(request.user):
        action = request.POST.get('action', 'create')
        if action == 'toggle':
            period = get_object_or_404(InventoryPeriod, pk=request.POST.get('period_id'))
            period.is_open = not period.is_open
            period.save(update_fields=['is_open', 'updated_at'])
            messages.success(request,
                             f'دوره {period.year} {"باز" if period.is_open else "بسته"} شد.')
        else:
            year = int(request.POST.get('year', '') or 0)
            if year:
                period, _ = InventoryPeriod.objects.get_or_create(
                    year=year, defaults={'title': request.POST.get('title', '')})
                messages.success(request, f'دوره {period.year} ثبت شد.')
        return redirect('inventory:periods')
    return render(request, 'inventory/period_list.html', {
        'periods': InventoryPeriod.objects.all(),
        'can_manage': _can_manage(request.user),
    })


def _excel_response(filename: str, headers: list[str], rows: list[list]) -> HttpResponse:
    """پاسخ Excel (openpyxl) — الگوی مشترک گزارش‌ها."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'گزارش'
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        sheet.append(row)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


@login_required
def balance_excel(request):
    if not feature_enabled('inventory.excel_reports'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.excel_reports'})
    warehouse = None
    if request.GET.get('warehouse'):
        warehouse = Warehouse.objects.filter(
            pk=request.GET.get('warehouse')).first()
    rows = [[item.code, item.name, item.unit,
             float(current_stock(item, warehouse)),
             float(item.min_stock), float(item.reorder_point)]
            for item in Item.objects.filter(is_active=True)]
    return _excel_response('balance.xlsx',
                           ['کد کالا', 'نام کالا', 'واحد', 'موجودی', 'حداقل', 'نقطه سفارش'],
                           rows)


@login_required
def movement_excel(request, pk):
    if not feature_enabled('inventory.excel_reports'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.excel_reports'})
    item = get_object_or_404(Item, pk=pk)
    report = movement_report(item)
    rows = [[str(r['move'].date), r['move'].reference,
             r['move'].get_kind_display(), r['move'].warehouse.name,
             float(r['move'].signed_quantity),
             float(r['move'].unit_price or 0), float(r['balance'])]
            for r in report['rows']]
    return _excel_response(f'movement-{item.code}.xlsx',
                           ['تاریخ', 'مرجع', 'نوع', 'انبار', 'تعداد', 'قیمت واحد', 'مانده'],
                           rows)


@login_required
def accounting_excel(request):
    """خروجی اسناد برای حسابداری (پل یکپارچه‌سازی)."""
    if not feature_enabled('inventory.excel_reports'):
        return render(request, 'core/feature_disabled.html',
                      {'feature_code': 'inventory.excel_reports'})
    from .services import export_accounting_rows
    date_from = _parse_jalali(request.GET.get('date_from', ''))
    date_to = _parse_jalali(request.GET.get('date_to', ''))
    rows = [[str(r['date']), r['number'], r['item_code'], r['item_name'],
             r['warehouse'], r['kind'], float(r['quantity']),
             float(r['unit_price']), float(r['value']),
             r['reference'], r['note']]
            for r in export_accounting_rows(date_from, date_to)]
    return _excel_response('accounting.xlsx',
                           ['تاریخ', 'شماره سند', 'کد کالا', 'نام کالا', 'انبار', 'نوع',
                            'تعداد', 'قیمت واحد', 'ارزش', 'مرجع', 'شرح'],
                           rows)


@login_required
def count_template_excel(request):
    """قالب اکسل ورود شمارش انبارگردانی."""
    return _excel_response('count-template.xlsx',
                           ['کد کالا', 'تعداد شمارش‌شده'],
                           [[item.code, ''] for item in Item.objects.filter(is_active=True)])


@login_required
def dashboard(request):
    """داشبورد انبار — موجودی، هشدارها و آخرین اسناد."""
    warehouses = _warehouses(request)
    low = low_stock_items()
    recent_docs = StockDocument.objects.select_related(
        'warehouse', 'created_by')[:8]
    counts = StockCount.objects.select_related('warehouse')[:5]
    total_items = Item.objects.filter(is_active=True).count()
    show_value = _can_view_value(request.user)
    value_data = inventory_value() if show_value else {'total': 0}
    return render(request, 'inventory/dashboard.html', {
        'warehouses': warehouses, 'low': low, 'recent_docs': recent_docs,
        'counts': counts, 'total_items': total_items,
        'total_value': value_data['total'], 'show_value': show_value,
        'can_manage': _can_manage(request.user),
    })
