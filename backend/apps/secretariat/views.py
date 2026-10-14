"""نماهای دبیرخانه و کارتابل — فاز ۱ (رابط مصوب «آرام و شفاف»)."""
from __future__ import annotations

import io
from datetime import date, timedelta
from urllib.parse import quote

import jdatetime
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.core.persian import to_persian_digits
from .forms import (AnswerActionForm, CopyForm, DelegationForm, IncomingLetterForm,
                    InternalLetterForm, LetterNoteForm, LetterRelationForm,
                    OutgoingDraftForm, ReferralActionForm)
from .models import (Attachment, Copy, Letter, LetterApproval, LetterEvent,
                     LetterNote, LetterRelation, LetterTemplate, Office, Referral)
from .services import (apply_confidentiality, current_jalali_year, register_letter,
                       visible_letters)

CLOSED_STATUSES = (Letter.Status.CLOSED, Letter.Status.ARCHIVED)
STEP_INDEX = {
    'draft': 0, 'registered': 1, 'in_review': 1, 'referred': 2,
    'in_progress': 2, 'waiting_user': 2, 'answered': 3,
    'signed': 3, 'sent': 4, 'closed': 4, 'archived': 4,
}


def deadline_gregorian(letter):
    d = letter.deadline
    if d is None:
        return None
    return d.togregorian() if hasattr(d, 'togregorian') else d


def _parse_jalali_input(raw: str):
    """رشته ورودی شمسی → jdatetime.date یا None (بدون خطا)."""
    from django import forms as dj_forms
    from .forms import JalaliDateField
    if not raw or not raw.strip():
        return None
    try:
        return JalaliDateField().clean(raw)
    except dj_forms.ValidationError:
        return None


def _notify(user, title, body='', url=''):
    """اعلان یکپارچه — واگذاری به سرویس مرکزی چندکاناله (فاز ۲).

    کانال‌ها (درون‌سامانه‌ای/ایمیل/پیامک) در apps.core.services تصمیم
    گرفته می‌شوند؛ این تابع فقط سازگاری با امضای قدیمی را نگه می‌دارد.
    """
    if user is None:
        return None
    from apps.core.services import send_notification
    return send_notification(user, title, body, url)


def _form_errors_message(form):
    for field_errors in form.errors.values():
        for error in field_errors:
            return str(error)
    for error in form.non_field_errors():
        return str(error)
    return 'خطا در فرم.'


# ---------- ثبت نامه ----------

@login_required
def register_letter_view(request):
    can_incoming = request.user.is_superuser or request.user.has_perm('secretariat.manage_secretariat')
    incoming = IncomingLetterForm(prefix='in') if can_incoming else None
    outgoing = OutgoingDraftForm(prefix='out')

    if request.method == 'POST':
        which = request.POST.get('form')
        if which == 'incoming':
            if not can_incoming:
                messages.error(request, 'مجوز ثبت نامه وارده ندارید.')
                return redirect('secretariat:register')
            incoming = IncomingLetterForm(request.POST, request.FILES, prefix='in')
            outgoing = OutgoingDraftForm(prefix='out')
            if incoming.is_valid():
                letter = incoming.save(request.user)
                messages.success(request, f'نامه وارده با شماره {letter.number_display} ثبت شد.')
                for name in incoming.invalid_attachments:
                    messages.warning(request, f'پیوست «{name}» به دلیل فرمت/حجم پذیرفته نشد.')
                return redirect('secretariat:letter_detail', pk=letter.pk)
        elif which == 'outgoing':
            outgoing = OutgoingDraftForm(request.POST, prefix='out')
            incoming = IncomingLetterForm(prefix='in') if can_incoming else None
            if outgoing.is_valid():
                letter = outgoing.save(request.user)
                messages.success(request, f'پیش‌نویس صادره با شماره {letter.number_display} ایجاد شد.')
                return redirect('secretariat:letter_detail', pk=letter.pk)
    return render(request, 'secretariat/register.html', {
        'incoming_form': incoming,
        'outgoing_form': outgoing,
        'can_incoming': can_incoming,
    })


# ---------- دفترچه‌ها ----------

@login_required
def letter_list(request):
    qs = visible_letters(request.user).select_related('office', 'created_by')
    base_ids = qs.values_list('pk', flat=True)

    letter_type = request.GET.get('type', '')
    if letter_type in dict(Letter.LetterType.choices):
        qs = qs.filter(letter_type=letter_type)

    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(
            Q(subject__icontains=q) | Q(summary__icontains=q)
            | Q(sender_name__icontains=q) | Q(sender_org__icontains=q))

    status_f = request.GET.get('status', '')
    if status_f in dict(Letter.Status.choices):
        qs = qs.filter(status=status_f)
    priority_f = request.GET.get('priority', '')
    if priority_f in dict(Letter.Priority.choices):
        qs = qs.filter(priority=priority_f)
    conf_f = request.GET.get('conf', '')
    if conf_f in dict(Letter.Confidentiality.choices):
        qs = qs.filter(confidentiality=conf_f)
    office_f = request.GET.get('office', '')
    if office_f.isdigit():
        qs = qs.filter(office_id=int(office_f))

    # شمارش زبانه‌ها روی نتیجه جست‌وجو (بدون فیلتر نوع)
    counted = visible_letters(request.user)
    if q:
        counted = counted.filter(
            Q(subject__icontains=q) | Q(summary__icontains=q)
            | Q(sender_name__icontains=q) | Q(sender_org__icontains=q))
    counts = {
        'all': counted.count(),
        'incoming': counted.filter(letter_type='incoming').count(),
        'outgoing': counted.filter(letter_type='outgoing').count(),
        'internal': counted.filter(letter_type='internal').count(),
    }

    date_from = _parse_jalali_input(request.GET.get('date_from', ''))
    date_to = _parse_jalali_input(request.GET.get('date_to', ''))
    if date_from:
        qs = qs.filter(letter_date__gte=date_from.togregorian())
    if date_to:
        qs = qs.filter(letter_date__lte=date_to.togregorian())

    if request.GET.get('export') == 'xlsx':
        return _export_xlsx(qs.order_by('-jalali_year', '-serial')[:5000])

    paginator = Paginator(qs.order_by('-jalali_year', '-serial'), 20)
    page = paginator.get_page(request.GET.get('page'))
    params = request.GET.copy()
    params.pop('page', None)
    return render(request, 'secretariat/letter_list.html', {
        'page_obj': page,
        'letters': page.object_list,
        'counts': counts,
        'query_base': params.urlencode(),
        'offices': Office.objects.filter(is_active=True),
        'status_choices': Letter.Status.choices,
        'priority_choices': Letter.Priority.choices,
        'conf_choices': Letter.Confidentiality.choices,
        'cur': {'type': letter_type, 'q': q, 'status': status_f, 'priority': priority_f,
                'conf': conf_f, 'office': office_f,
                'date_from': request.GET.get('date_from', ''),
                'date_to': request.GET.get('date_to', '')},
    })


def _export_xlsx(letters):
    wb = Workbook()
    ws = wb.active
    ws.title = 'دفتر نامه'
    ws.sheet_view.rightToLeft = True
    headers = ['شماره', 'نوع', 'موضوع', 'فرستنده', 'تاریخ نامه', 'وضعیت', 'فوریت', 'محرمانگی', 'مهلت', 'دبیرخانه']
    header_fill = PatternFill('solid', fgColor='16365F')
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color='FFFFFF', name='Vazirmatn')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    type_labels = dict(Letter.LetterType.choices)
    for row, letter in enumerate(letters, 2):
        j_letter_date = letter.letter_date
        values = [
            letter.number_display,
            type_labels.get(letter.letter_type, ''),
            letter.subject,
            letter.sender_org or letter.sender_name or '—',
            to_persian_digits(j_letter_date.strftime('%Y/%m/%d')) if j_letter_date else '',
            letter.get_status_display(),
            letter.get_priority_display(),
            letter.get_confidentiality_display(),
            to_persian_digits(letter.deadline.strftime('%Y/%m/%d')) if letter.deadline else '',
            letter.office.name,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
    widths = [16, 10, 46, 24, 14, 14, 10, 12, 14, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = quote('daftar-name-ha.xlsx')
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response


# ---------- کارت نامه ----------

@login_required
def letter_detail(request, pk):
    letter = get_object_or_404(
        visible_letters(request.user).select_related('office', 'created_by'), pk=pk)
    user = request.user
    can_refer = user.is_superuser or user.has_perm('secretariat.refer_letter')
    can_answer = user.is_superuser or user.has_perm('secretariat.answer_letter')
    can_close = user.is_superuser or user.has_perm('secretariat.close_letter')
    is_open = letter.status not in CLOSED_STATUSES

    referrals = letter.referrals.select_related('to_user', 'to_department', 'from_user', 'parent')
    can_approve = (letter.letter_type == 'outgoing'
                   and letter.status == Letter.Status.IN_REVIEW
                   and (user.is_superuser or user.has_perm('secretariat.approve_letter')))
    can_submit = (letter.letter_type == 'outgoing' and letter.status == Letter.Status.DRAFT
                  and (letter.created_by_id == user.pk or user.is_superuser
                       or user.has_perm('secretariat.manage_secretariat')))
    can_send = (letter.letter_type == 'outgoing' and letter.status == Letter.Status.SIGNED
                and (letter.created_by_id == user.pk or user.is_superuser
                     or user.has_perm('secretariat.manage_secretariat')
                     or user.has_perm('secretariat.approve_letter')))
    # فاز ۲ — کارت‌های گردش کار، بایگانی و وظایف روی صفحه نامه
    from django.contrib.contenttypes.models import ContentType
    from apps.archive.models import ArchiveItem
    from apps.workflow.services import instance_for
    ct_letter = ContentType.objects.get_for_model(Letter)
    archive_item = ArchiveItem.objects.filter(
        content_type=ct_letter, object_id=letter.pk,
    ).select_related('folder', 'folder__box').first()
    return render(request, 'secretariat/letter_detail.html', {
        'letter': letter,
        'wf_instance': instance_for(letter),
        'can_bind_workflow': user.is_superuser
                             or user.has_perm('secretariat.manage_secretariat'),
        'archive_item': archive_item,
        'can_archive': letter.status in (Letter.Status.CLOSED, Letter.Status.ARCHIVED)
                       and (user.is_superuser or user.has_perm('archive.manage_archive')),
        'letter_tasks': letter.tasks.select_related('responsible')[:10],
        'can_task_from_letter': can_refer and is_open,
        'approvals': letter.approvals.select_related('approver'),
        'can_submit': can_submit,
        'can_approve': can_approve,
        'can_send': can_send,
        'referrals': referrals,
        'copies': letter.copies.select_related('to_user', 'to_department'),
        'attachments': letter.attachments.all(),
        'events': letter.events.select_related('actor')[:50],
        'step_index': STEP_INDEX.get(letter.status, 0),
        'can_refer': can_refer and is_open,
        'can_answer': can_answer and is_open,
        'can_close': can_close and is_open,
        'referral_form': ReferralActionForm() if can_refer and is_open else None,
        'answer_form': AnswerActionForm() if can_answer and is_open else None,
        'copy_form': CopyForm() if can_refer and is_open else None,
        'deadline_g': deadline_gregorian(letter),
        'relations': LetterRelation.objects.filter(
            Q(letter=letter) | Q(related=letter)
        ).select_related('letter', 'related'),
        'notes': letter.notes.select_related('author'),
        'relation_form': LetterRelationForm(),
        'note_form': LetterNoteForm(),
    })


@login_required
@require_POST
def letter_refer(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('secretariat.refer_letter')):
        messages.error(request, 'مجوز ارجاع نامه ندارید.')
        return redirect('secretariat:letter_detail', pk=pk)
    if letter.status in CLOSED_STATUSES:
        messages.error(request, 'نامه بسته/بایگانی شده قابل ارجاع نیست.')
        return redirect('secretariat:letter_detail', pk=pk)
    form = ReferralActionForm(request.POST)
    if not form.is_valid():
        messages.error(request, _form_errors_message(form))
        return redirect('secretariat:letter_detail', pk=pk)
    d = form.cleaned_data
    parent_id = request.POST.get('parent')
    parent = letter.referrals.filter(pk=parent_id).first() if parent_id else None
    referral = Referral.objects.create(
        letter=letter, parent=parent, from_user=request.user,
        to_user=d['to_user'], to_department=d['to_department'],
        note=d['note'], deadline=d['deadline'])
    old = letter.status
    letter.status = Letter.Status.REFERRED
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.REFERRED,
        from_status=old, to_status=letter.status,
        note=(d['note'] or '')[:300])
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    _notify(d['to_user'], f'ارجاع نامه {letter.number_display} — {letter.subject[:60]}',
            d['note'], url)
    messages.success(request, 'ارجاع ثبت شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_answer(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('secretariat.answer_letter')):
        messages.error(request, 'مجوز ثبت پاسخ ندارید.')
        return redirect('secretariat:letter_detail', pk=pk)
    if letter.status in CLOSED_STATUSES:
        messages.error(request, 'نامه بسته/بایگانی شده قابل پاسخ نیست.')
        return redirect('secretariat:letter_detail', pk=pk)
    form = AnswerActionForm(request.POST)
    if not form.is_valid():
        messages.error(request, _form_errors_message(form))
        return redirect('secretariat:letter_detail', pk=pk)
    note = form.cleaned_data['note']
    old = letter.status
    letter.status = Letter.Status.ANSWERED
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.ANSWERED,
        from_status=old, to_status=letter.status, note=note[:500])
    Referral.objects.filter(
        letter=letter, to_user=request.user, status=Referral.Status.PENDING,
    ).update(status=Referral.Status.ACTED, acted_at=timezone.now())
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    recipients = {letter.created_by_id}
    recipients.update(letter.referrals.values_list('from_user_id', flat=True))
    from django.contrib.auth import get_user_model
    User = get_user_model()
    for uid in filter(None, recipients):
        if uid != request.user.pk:
            _notify(User.objects.filter(pk=uid).first(),
                    f'پاسخ نامه {letter.number_display} ثبت شد', note, url)
    messages.success(request, 'پاسخ ثبت شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_attach(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    f = request.FILES.get('file')
    if not f:
        messages.error(request, 'فایلی انتخاب نشد.')
        return redirect('secretariat:letter_detail', pk=pk)
    attachment = Attachment(letter=letter, file=f, uploaded_by=request.user)
    try:
        attachment.full_clean()
    except Exception as exc:  # ValidationError
        message = exc.messages[0] if hasattr(exc, 'messages') and exc.messages else 'پیوست پذیرفته نشد.'
        messages.error(request, message)
        return redirect('secretariat:letter_detail', pk=pk)
    attachment.save()
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.ATTACHED,
        note=attachment.original_name)
    messages.success(request, f'پیوست «{attachment.original_name}» افزوده شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_add_copy(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    form = CopyForm(request.POST)
    if not form.is_valid():
        messages.error(request, _form_errors_message(form))
        return redirect('secretariat:letter_detail', pk=pk)
    d = form.cleaned_data
    Copy.objects.create(
        letter=letter, to_user=d['to_user'], to_department=d['to_department'], note=d['note'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.COPIED,
        note=d['note'][:300])
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    _notify(d['to_user'], f'رونوشت نامه {letter.number_display}', d['note'], url)
    messages.success(request, 'رونوشت افزوده شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_close(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('secretariat.close_letter')):
        messages.error(request, 'مجوز بستن/بایگانی نامه را ندارید.')
        return redirect('secretariat:letter_detail', pk=pk)
    old = letter.status
    letter.status = Letter.Status.ARCHIVED
    letter.save(update_fields=['status', 'updated_at'])
    Referral.objects.filter(letter=letter, status=Referral.Status.PENDING).update(
        status=Referral.Status.CLOSED)
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.ARCHIVED,
        from_status=old, to_status=letter.status)
    messages.success(request, 'نامه بایگانی شد.')
    return redirect('secretariat:letter_detail', pk=pk)


# ---------- کارتابل ----------

@login_required
def kartable(request):
    user = request.user
    today = date.today()
    warning_days = today + timedelta(days=3)

    my_referrals_pending = Referral.objects.filter(
        to_user=user, status=Referral.Status.PENDING).select_related('letter', 'from_user', 'letter__office')
    waiting_ids = [r.letter_id for r in my_referrals_pending]
    waiting_letters = list(apply_confidentiality(
        Letter.objects.filter(pk__in=waiting_ids), user).select_related('office'))
    # نامه‌های داخلی تحویل‌شده به کارتابل این کاربر
    have_ids = {letter.pk for letter in waiting_letters}
    for letter in apply_confidentiality(
            Letter.objects.filter(to_user=user).exclude(
                status__in=CLOSED_STATUSES), user).select_related('office'):
        if letter.pk not in have_ids:
            waiting_letters.append(letter)
            have_ids.add(letter.pk)
    visible_ids = have_ids
    referral_map = {r.letter_id: r for r in my_referrals_pending if r.letter_id in visible_ids}

    in_progress = Referral.objects.filter(
        to_user=user, status=Referral.Status.ACTED,
    ).exclude(letter__status__in=CLOSED_STATUSES).select_related(
        'letter', 'letter__office').order_by('-acted_at')[:50]

    followups = Referral.objects.filter(
        from_user=user, status=Referral.Status.PENDING,
    ).select_related('letter', 'to_user', 'to_department', 'letter__office').order_by('-created_at')[:50]

    sent = Letter.objects.filter(created_by=user).exclude(
        status__in=CLOSED_STATUSES).select_related('office').order_by('-jalali_year', '-serial')[:50]

    archive = apply_confidentiality(
        Letter.objects.filter(
            Q(created_by=user) | Q(referrals__to_user=user) | Q(copies__to_user=user)
            | Q(to_user=user),
            status__in=CLOSED_STATUSES,
        ).distinct(), user).select_related('office')[:50]

    near = overdue = 0
    for letter in waiting_letters:
        g = deadline_gregorian(letter)
        referral = referral_map.get(letter.pk)
        if referral is not None and referral.deadline is not None:
            rg = (referral.deadline.togregorian()
                  if hasattr(referral.deadline, 'togregorian') else referral.deadline)
            g = min(g, rg) if g else rg
        if g is None:
            continue
        if g < today:
            overdue += 1
        elif g <= warning_days:
            near += 1

    answered_events = LetterEvent.objects.filter(
        letter__created_by=user, action=LetterEvent.Action.ANSWERED).select_related('letter')
    total_answered = answered_events.count()
    on_time = sum(
        1 for event in answered_events
        if deadline_gregorian(event.letter) and event.created_at.date() <= deadline_gregorian(event.letter))
    rate = round(on_time / total_answered * 100) if total_answered else None

    tab = request.GET.get('tab', 'waiting')
    tickets = _my_tickets(user, tab)
    return render(request, 'secretariat/kartable.html', {
        'tab': tab,
        'tickets': tickets,
        'waiting': waiting_letters,
        'referral_map': referral_map,
        'in_progress': [r.letter for r in in_progress],
        'followups': followups,
        'sent': sent,
        'archive': archive,
        'kpi': {
            'waiting': len(waiting_letters),
            'near': near,
            'overdue': overdue,
            'in_progress': len(in_progress),
            'rate': to_persian_digits(rate) if rate is not None else '—',
        },
    })


def _my_tickets(user, tab: str):
    """تیکت‌های کاربر از پل PHP (فقط‌خواندنی) — docs/14 بند ۷."""
    if tab != 'tickets':
        return []
    try:
        from apps.ticket_bridge.models import ExternalTicket
    except ImportError:
        return []
    usernames = list(
        user.integration_maps.values_list('external_username', flat=True)) or [user.username]
    return ExternalTicket.objects.filter(
        requester_username__in=usernames).order_by('-created_label')[:50]


# ---------- گردش صادره: بازبینی، تأیید/امضا، ارسال ----------

def _notify_approvers(letter):
    from django.contrib.auth import get_user_model
    from django.contrib.auth.models import Permission
    User = get_user_model()
    perm = Permission.objects.filter(
        codename='approve_letter', content_type__app_label='secretariat').first()
    if not perm:
        return
    approvers = User.objects.filter(
        Q(groups__permissions=perm) | Q(user_permissions=perm), is_active=True).distinct()
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    for approver in approvers:
        _notify(approver, f'نامه «{letter.subject[:50]}» در انتظار بازبینی/امضا است', '', url)


@login_required
@require_POST
def letter_submit(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    allowed = (letter.created_by_id == request.user.pk or request.user.is_superuser
               or request.user.has_perm('secretariat.manage_secretariat'))
    if letter.letter_type != 'outgoing' or letter.status != Letter.Status.DRAFT or not allowed:
        messages.error(request, 'ارسال به بازبینی فقط برای پیش‌نویس صادره خودتان ممکن است.')
        return redirect('secretariat:letter_detail', pk=pk)
    letter.status = Letter.Status.IN_REVIEW
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.STATUS_CHANGED,
        from_status=Letter.Status.DRAFT, to_status=letter.status, note='ارسال به بازبینی')
    _notify_approvers(letter)
    messages.success(request, 'نامه برای بازبینی/امضا ارسال شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_approve(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('secretariat.approve_letter')):
        messages.error(request, 'مجوز تأیید/امضا ندارید.')
        return redirect('secretariat:letter_detail', pk=pk)
    if letter.letter_type != 'outgoing' or letter.status != Letter.Status.IN_REVIEW:
        messages.error(request, 'فقط نامه صادره در بازبینی قابل تأیید است.')
        return redirect('secretariat:letter_detail', pk=pk)
    approval = LetterApproval.objects.create(
        letter=letter, step=letter.approvals.count() + 1, approver=request.user,
        approver_label=request.POST.get('label', '')[:150],
        decision=LetterApproval.Decision.APPROVED,
        note=request.POST.get('note', '')[:500],
        decided_at=timezone.now(),
        signature_image=request.FILES.get('signature') or None)
    letter.status = Letter.Status.SIGNED
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.STATUS_CHANGED,
        from_status=Letter.Status.IN_REVIEW, to_status=letter.status,
        note=f'تأیید/امضا — {approval.approver_label or "امضا"}')
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    _notify(letter.created_by, f'نامه {letter.number_display} امضا شد', approval.note, url)
    messages.success(request, 'نامه تأیید و امضا شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_reject(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('secretariat.approve_letter')):
        messages.error(request, 'مجوز بازبینی ندارید.')
        return redirect('secretariat:letter_detail', pk=pk)
    if letter.letter_type != 'outgoing' or letter.status != Letter.Status.IN_REVIEW:
        messages.error(request, 'فقط نامه صادره در بازبینی قابل رد است.')
        return redirect('secretariat:letter_detail', pk=pk)
    note = request.POST.get('note', '')[:500]
    LetterApproval.objects.create(
        letter=letter, step=letter.approvals.count() + 1, approver=request.user,
        approver_label=request.POST.get('label', '')[:150],
        decision=LetterApproval.Decision.REJECTED, note=note, decided_at=timezone.now())
    letter.status = Letter.Status.DRAFT
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.STATUS_CHANGED,
        from_status=Letter.Status.IN_REVIEW, to_status=letter.status,
        note=f'رد در بازبینی — {note}')
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    _notify(letter.created_by, f'نامه {letter.number_display} در بازبینی رد شد', note, url)
    messages.warning(request, 'نامه رد و به پیش‌نویس برگشت.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_send(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    allowed = (letter.created_by_id == request.user.pk or request.user.is_superuser
               or request.user.has_perm('secretariat.manage_secretariat')
               or request.user.has_perm('secretariat.approve_letter'))
    if letter.letter_type != 'outgoing' or letter.status != Letter.Status.SIGNED or not allowed:
        messages.error(request, 'ثبت ارسال فقط برای نامه امضاشده ممکن است.')
        return redirect('secretariat:letter_detail', pk=pk)
    letter.status = Letter.Status.SENT
    letter.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.STATUS_CHANGED,
        from_status=Letter.Status.SIGNED, to_status=letter.status,
        note='ارسال/ثبت در دفتر صادره')
    url = reverse('secretariat:letter_detail', args=[letter.pk])
    for copy in letter.copies.select_related('to_user'):
        _notify(copy.to_user, f'رونوشت نامه ارسال‌شده {letter.number_display}', letter.subject, url)
    messages.success(request, 'نامه ارسال و در دفتر صادره ثبت شد.')
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
def letter_pdf(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    from .pdf import build_letter_pdf
    data = build_letter_pdf(letter)
    response = HttpResponse(data, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="letter-{letter.pk}.pdf"'
    return response


@login_required
def letter_print(request, pk):
    letter = get_object_or_404(
        visible_letters(request.user).select_related(
            'office', 'created_by', 'to_user', 'to_department'), pk=pk)
    from .models import LetterRelation
    from django.db.models import Q
    return render(request, 'secretariat/letter_print.html', {
        'letter': letter,
        'attachments': letter.attachments.all(),
        'copies': letter.copies.select_related('to_user', 'to_department'),
        'related_in': LetterRelation.objects.filter(
            related=letter).select_related('letter'),
        'related_out': LetterRelation.objects.filter(
            letter=letter).select_related('related'),
    })


# ---------- نامه داخلی، نامه‌های مرتبط و یادداشت‌ها ----------

@login_required
def letter_internal(request):
    """ثبت و ارسال نامه داخلی — تحویل مستقیم به کارتابل گیرنده."""
    form = InternalLetterForm(request.POST or None, request.FILES or None,
                              prefix='int')
    if request.method == 'POST':
        if form.is_valid():
            letter = form.save(request.user)
            messages.success(
                request, f'نامه داخلی {letter.number_display} ثبت و ارسال شد.')
            for name in form.invalid_attachments:
                messages.warning(request, f'پیوست «{name}» پذیرفته نشد.')
            return redirect('secretariat:letter_detail', pk=letter.pk)
        messages.error(request, _form_errors_message(form))
    templates = LetterTemplate.objects.all()
    return render(request, 'secretariat/letter_internal.html', {
        'form': form, 'templates': templates,
    })


@login_required
@require_POST
def letter_duplicate(request, pk):
    """نامه مشابه — کپی متن و مشخصات به‌عنوان پیش‌نویس تازه (الگوی دیدگاه)."""
    source = get_object_or_404(visible_letters(request.user), pk=pk)
    office = source.office
    letter = register_letter(
        office, source.letter_type, source.subject, created_by=request.user,
        status=Letter.Status.DRAFT,
        summary=source.summary, body=source.body,
        to_user=source.to_user, to_department=source.to_department,
        priority=source.priority, confidentiality=source.confidentiality,
        delivery_method=source.delivery_method)
    LetterRelation.objects.get_or_create(
        letter=letter, related=source, kind=LetterRelation.Kind.FOLLOW)
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.REGISTERED,
        note=f'نامه مشابه {source.number_display} ساخته شد')
    messages.success(request, f'نامه مشابه با شماره {letter.number_display} ساخته شد.')
    return redirect('secretariat:letter_detail', pk=letter.pk)


@login_required
@require_POST
def letter_relation_add(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    form = LetterRelationForm(request.POST)
    if form.is_valid():
        related = Letter.objects.get(pk=form.cleaned_data['related_pk'])
        if related.pk == letter.pk:
            messages.error(request, 'نامه نمی‌تواند به خودش مرتبط باشد.')
        else:
            LetterRelation.objects.get_or_create(
                letter=letter, related=related,
                kind=form.cleaned_data['kind'])
            LetterEvent.objects.create(
                letter=letter, actor=request.user, action=LetterEvent.Action.NOTE,
                note=f'نامه مرتبط ({form.cleaned_data["kind"]}): '
                     f'{related.number_display}')
            messages.success(request, 'نامه مرتبط ثبت شد.')
    else:
        messages.error(request, _form_errors_message(form))
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def letter_note_add(request, pk):
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    form = LetterNoteForm(request.POST)
    if form.is_valid():
        LetterNote.objects.create(
            letter=letter, author=request.user, text=form.cleaned_data['text'][:2000])
        messages.success(request, 'یادداشت ثبت شد.')
    else:
        messages.error(request, _form_errors_message(form))
    return redirect('secretariat:letter_detail', pk=pk)


@login_required
@require_POST
def refer_cancel(request, pk):
    """بازپس‌گیری ارجاع — فقط ارجاع‌دهنده، پیش از اقدام گیرنده."""
    letter = get_object_or_404(visible_letters(request.user), pk=pk)
    referral = get_object_or_404(
        Referral, pk=request.POST.get('referral_id'),
        letter=letter, from_user=request.user,
        status=Referral.Status.PENDING)
    referral.status = Referral.Status.CANCELLED
    referral.save(update_fields=['status', 'updated_at'])
    LetterEvent.objects.create(
        letter=letter, actor=request.user, action=LetterEvent.Action.NOTE,
        note=f'ارجاع به {referral.to_user or referral.to_department} بازپس گرفته شد')
    messages.success(request, 'ارجاع بازپس گرفته شد.')
    return redirect('secretariat:letter_detail', pk=pk)

@login_required
def delegations(request):
    form = DelegationForm(request.POST or None, from_user=request.user)
    if request.method == 'POST' and form.is_valid():
        delegation = form.save()
        messages.success(request, f'جانشینی برای {delegation.to_user.get_full_name() or delegation.to_user.username} ثبت شد.')
        _notify(delegation.to_user, 'شما جانشین کارتابل شدید',
                f'از {delegation.start_date} تا {delegation.end_date}',
                reverse('secretariat:kartable'))
        return redirect('secretariat:delegations')
    return render(request, 'secretariat/delegations.html', {
        'form': form,
        'given': request.user.delegations_given.select_related('to_user'),
        'received': request.user.delegations_received.select_related('from_user'),
    })


@login_required
@require_POST
def delegation_deactivate(request, pk):
    delegation = get_object_or_404(request.user.delegations_given, pk=pk)
    delegation.is_active = False
    delegation.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'جانشینی غیرفعال شد.')
    return redirect('secretariat:delegations')
